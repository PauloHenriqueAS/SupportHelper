"""Extrai texto dos formatos aceitos. Imagens sao ignoradas."""

from dataclasses import dataclass
from pathlib import Path
from typing import Optional

SUPPORTED_EXTENSIONS = {".txt", ".md", ".docx", ".odt", ".pdf"}


@dataclass
class TextBlock:
    text: str
    page: Optional[int] = None


def _read_text_file(path: Path) -> list[TextBlock]:
    for encoding in ("utf-8", "utf-8-sig", "cp1252", "latin-1"):
        try:
            return [TextBlock(path.read_text(encoding=encoding))]
        except UnicodeDecodeError:
            continue
    return [TextBlock(path.read_bytes().decode("utf-8", errors="ignore"))]


def _read_docx(path: Path) -> list[TextBlock]:
    import docx  # python-docx

    document = docx.Document(str(path))
    parts: list[str] = []
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            parts.append(paragraph.text)
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            line = " | ".join(c for c in cells if c)
            if line:
                parts.append(line)
    return [TextBlock("\n".join(parts))]


def _read_odt(path: Path) -> list[TextBlock]:
    from odf import teletype
    from odf.opendocument import load
    from odf.text import H, P

    document = load(str(path))
    parts: list[str] = []
    for element in document.getElementsByType(H) + document.getElementsByType(P):
        text = teletype.extractText(element).strip()
        if text:
            parts.append(text)
    return [TextBlock("\n".join(parts))]


def _read_pdf(path: Path) -> list[TextBlock]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    blocks: list[TextBlock] = []
    for index, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            blocks.append(TextBlock(text, page=index))
    return blocks


_READERS = {
    ".txt": _read_text_file,
    ".md": _read_text_file,
    ".docx": _read_docx,
    ".odt": _read_odt,
    ".pdf": _read_pdf,
}


def is_supported(path: Path) -> bool:
    return path.suffix.lower() in SUPPORTED_EXTENSIONS


def extract_text(path: Path) -> list[TextBlock]:
    """Retorna blocos de texto do arquivo. Lanca ValueError para formatos nao suportados."""
    reader = _READERS.get(path.suffix.lower())
    if reader is None:
        raise ValueError(f"Formato nao suportado: {path.suffix}")
    blocks = reader(path)
    return [b for b in blocks if b.text.strip()]
